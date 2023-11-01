import numpy as np
from numpy import random
import docx
from openpyxl import Workbook

N = int(input())
M = int(input())
matrix = []
for i in range(N):
  row = [round(random.uniform(0, 1), 1) for _ in range(M)]
  matrix.append(row)
matrix = np.array(matrix)

doc = docx.Document()
table = doc.add_table(rows=N, cols=M)
for i in range(N):
  row_cells = table.rows[i].cells
  for j in range(M):
    row_cells[j].text = str(matrix[i][j])
doc.save("matrix.docx")


wb = Workbook()
ws = wb.active
for i in range(N):
  for j in range(M):
    ws.cell(row=i + 1, column=j + 1, value=matrix[i][j])
wb.save("matrix.xlsx")